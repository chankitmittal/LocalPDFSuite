"""
Batch Form Data Extraction Engine
Extracts interactive AcroForm field data across multiple PDF files and exports
to enterprise-grade formatted Excel (.xlsx via openpyxl) and CSV (.csv with UTF-8 BOM).

Features:
- Schema union auto-discovery across heterogeneous form templates.
- Data coercion: text, boolean/checkbox, dropdown choices, signatures.
- High-aesthetic Excel workbook with frozen header, auto-filter, zebra rows, auto column widths.
- Companion Schema Summary sheet with fill-rates and field-type telemetry.
- RFC 4180 compliant CSV with UTF-8 BOM for flawless Excel opening.
- 100% offline local processing.
"""

import os
import csv
import logging
from typing import List, Dict, Any, Optional
from pypdf import PdfReader
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


def extract_form_fields_from_pdf(pdf_path: str) -> Dict[str, Any]:
    """
    Extracts all AcroForm field keys and coerced values from a single PDF document.
    Returns:
        {
            "filename": str,
            "filepath": str,
            "page_count": int,
            "file_size": int,
            "fields": {field_name: {"value": val, "type": type_str, "raw": raw_val}},
            "field_count": int,
            "error": Optional[str]
        }
    """
    res = {
        "filename": os.path.basename(pdf_path),
        "filepath": pdf_path,
        "page_count": 0,
        "file_size": 0,
        "fields": {},
        "field_count": 0,
        "error": None,
    }

    if not os.path.exists(pdf_path):
        res["error"] = f"File not found: {pdf_path}"
        return res

    try:
        res["file_size"] = os.path.getsize(pdf_path)
        with open(pdf_path, "rb") as f_in:
            reader = PdfReader(f_in)
            res["page_count"] = len(reader.pages)
            raw_fields = reader.get_fields() or {}

        for name, data in raw_fields.items():
            if not isinstance(data, dict):
                continue

            ft = str(data.get("/FT", ""))
            raw_val = data.get("/V", "")

            field_type = "text"
            if ft == "/Btn":
                field_type = "checkbox"
            elif ft == "/Ch":
                field_type = "choice"
            elif ft == "/Sig":
                field_type = "signature"

            # Coerce value
            coerced_val = ""
            if raw_val is not None:
                if isinstance(raw_val, bool):
                    coerced_val = "Yes" if raw_val else "No"
                else:
                    s_val = str(raw_val).lstrip("/")
                    if field_type == "checkbox":
                        if s_val.lower() in ("yes", "on", "true", "1", "checked"):
                            coerced_val = "Yes"
                        elif s_val.lower() in ("off", "no", "false", "0", ""):
                            coerced_val = "No"
                        else:
                            coerced_val = s_val
                    else:
                        coerced_val = s_val.strip()
            elif field_type == "checkbox":
                coerced_val = "No"

            res["fields"][name] = {
                "value": coerced_val,
                "type": field_type,
                "raw": str(raw_val) if raw_val is not None else "",
            }

        res["field_count"] = len(res["fields"])
    except Exception as e:
        logger.error(f"Error reading form fields from {pdf_path}: {e}")
        res["error"] = str(e)

    return res


def batch_extract_forms_to_spreadsheet(
    pdf_paths: List[str],
    output_dir: str,
    base_filename: str = "batch_form_export",
) -> Dict[str, Any]:
    """
    Extracts AcroForm field data across all provided PDF paths, builds a unified
    schema, and writes formatted .xlsx and .csv files.

    Returns summary details including file paths, column schema, and sample rows.
    """
    os.makedirs(output_dir, exist_ok=True)

    extracted_records: List[Dict[str, Any]] = []
    field_occurrences: Dict[str, int] = {}
    field_types: Dict[str, str] = {}
    field_non_empty_counts: Dict[str, int] = {}

    for path in pdf_paths:
        rec = extract_form_fields_from_pdf(path)
        extracted_records.append(rec)
        for fname, fdata in rec["fields"].items():
            field_occurrences[fname] = field_occurrences.get(fname, 0) + 1
            if fname not in field_types:
                field_types[fname] = fdata["type"]
            val = fdata["value"]
            if val and val != "No":
                field_non_empty_counts[fname] = field_non_empty_counts.get(fname, 0) + 1

    total_docs = len(pdf_paths)

    # Sort fields: most frequently occurring first, then alphabetically
    sorted_field_names = sorted(
        field_occurrences.keys(),
        key=lambda k: (-field_occurrences[k], k.lower())
    )

    metadata_headers = ["Source File", "Pages", "File Size (KB)", "Fields Extracted", "Status"]
    all_headers = metadata_headers + sorted_field_names

    # Prepare rows
    data_rows: List[List[Any]] = []
    for rec in extracted_records:
        size_kb = round(rec["file_size"] / 1024, 1) if rec["file_size"] else 0
        status = "Error" if rec["error"] else ("Populated" if rec["field_count"] > 0 else "No Form Fields")
        meta_values = [rec["filename"], rec["page_count"], size_kb, rec["field_count"], status]

        field_vals = []
        for fn in sorted_field_names:
            finfo = rec["fields"].get(fn)
            val = finfo["value"] if finfo else ""
            field_vals.append(val)

        data_rows.append(meta_values + field_vals)

    # 1. Generate Formatted Excel (.xlsx)
    xlsx_path = os.path.join(output_dir, f"{base_filename}.xlsx")
    wb = openpyxl.Workbook()

    # Sheet 1: Extracted Data
    ws_data = wb.active
    ws_data.title = "Form Data"

    # Header styling
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    meta_fill_even = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    meta_fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    row_fill_even = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    row_fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # Write Headers
    ws_data.append(all_headers)
    ws_data.row_dimensions[1].height = 28

    for col_idx in range(1, len(all_headers) + 1):
        cell = ws_data.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border

    # Write Data
    for r_idx, row_values in enumerate(data_rows, start=2):
        ws_data.append(row_values)
        ws_data.row_dimensions[r_idx].height = 20
        is_even = (r_idx % 2 == 0)

        for c_idx, val in enumerate(row_values, start=1):
            cell = ws_data.cell(row=r_idx, column=c_idx)
            cell.border = thin_border
            cell.font = Font(name="Segoe UI", size=10)

            # Metadata column tinting vs data column tinting
            if c_idx <= len(metadata_headers):
                cell.fill = meta_fill_even if is_even else meta_fill_odd
                if c_idx in (2, 3, 4):  # Pages, Size, Fields
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif c_idx == 5:  # Status
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.fill = row_fill_even if is_even else row_fill_odd
                # Center boolean / Yes-No values
                if str(val).lower() in ("yes", "no"):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

    # Freeze Header Row and Enable Auto-Filter
    ws_data.freeze_panes = "A2"
    if len(all_headers) > 0 and len(data_rows) > 0:
        max_col_letter = get_column_letter(len(all_headers))
        ws_data.auto_filter.ref = f"A1:{max_col_letter}{len(data_rows) + 1}"

    # Auto-adjust column widths based on longest string
    for col in ws_data.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        # Bounded between 12 and 45 characters
        ws_data.column_dimensions[col_letter].width = max(12, min(max_len + 3, 45))

    # Sheet 2: Schema & Fill Rate Summary
    ws_summary = wb.create_sheet(title="Field Schema & Stats")
    ws_summary.append(["Field Name", "Field Type", "Documents Present", "Present %", "Populated Count", "Fill Rate %"])
    ws_summary.row_dimensions[1].height = 26

    summary_header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    for c_idx in range(1, 7):
        cell = ws_summary.cell(row=1, column=c_idx)
        cell.fill = summary_header_fill
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    schema_info_list = []
    for s_idx, fname in enumerate(sorted_field_names, start=2):
        docs_present = field_occurrences.get(fname, 0)
        docs_present_pct = round((docs_present / total_docs * 100), 1) if total_docs > 0 else 0
        populated = field_non_empty_counts.get(fname, 0)
        fill_rate_pct = round((populated / total_docs * 100), 1) if total_docs > 0 else 0
        ftype = field_types.get(fname, "text")

        ws_summary.append([fname, ftype, docs_present, f"{docs_present_pct}%", populated, f"{fill_rate_pct}%"])
        ws_summary.row_dimensions[s_idx].height = 19
        for c in range(1, 7):
            c_cell = ws_summary.cell(row=s_idx, column=c)
            c_cell.font = Font(name="Segoe UI", size=9.5)
            c_cell.border = thin_border
            if c in (3, 4, 5, 6):
                c_cell.alignment = Alignment(horizontal="center", vertical="center")

        schema_info_list.append({
            "field_name": fname,
            "type": ftype,
            "documents_present": docs_present,
            "present_pct": docs_present_pct,
            "populated_count": populated,
            "fill_rate_pct": fill_rate_pct,
        })

    ws_summary.freeze_panes = "A2"
    for col in ws_summary.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws_summary.column_dimensions[col_letter].width = max(14, min(max_len + 3, 40))

    wb.save(xlsx_path)

    # 2. Generate RFC 4180 Compliant CSV with UTF-8 BOM
    csv_path = os.path.join(output_dir, f"{base_filename}.csv")
    with open(csv_path, mode="w", newline="", encoding="utf-8-sig") as csv_file:
        writer = csv.writer(csv_file, quoting=csv.QUOTE_MINIMAL)
        writer.writerow(all_headers)
        for row in data_rows:
            writer.writerow(row)

    # Preview rows for UI display (first 10 rows, capped)
    preview_rows = []
    for r in data_rows[:10]:
        preview_rows.append([str(x) for x in r])

    return {
        "status": "success",
        "total_documents": total_docs,
        "total_fields": len(sorted_field_names),
        "headers": all_headers,
        "metadata_column_count": len(metadata_headers),
        "schema": schema_info_list,
        "preview_rows": preview_rows,
        "total_rows": len(data_rows),
        "xlsx_path": xlsx_path,
        "xlsx_filename": os.path.basename(xlsx_path),
        "xlsx_size": os.path.getsize(xlsx_path),
        "csv_path": csv_path,
        "csv_filename": os.path.basename(csv_path),
        "csv_size": os.path.getsize(csv_path),
    }
