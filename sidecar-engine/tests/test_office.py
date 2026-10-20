import os
import openpyxl
from docx import Document
from pptx import Presentation
from core.pdf_to_office import (
    pdf_to_excel,
    pdf_to_word,
    pdf_to_powerpoint,
)


def test_pdf_to_excel(sample_pdf, temp_dir):
    out_xlsx = os.path.join(temp_dir, "output.xlsx")
    res = pdf_to_excel(sample_pdf, out_xlsx)

    assert res["status"] == "success"
    assert os.path.exists(out_xlsx)

    wb = openpyxl.load_workbook(out_xlsx)
    assert len(wb.sheetnames) == 3


def test_pdf_to_word(sample_pdf, temp_dir):
    out_docx = os.path.join(temp_dir, "output.docx")
    res = pdf_to_word(sample_pdf, out_docx)

    assert res["status"] == "success"
    assert os.path.exists(out_docx)

    doc = Document(out_docx)
    # Check that paragraphs or headings were reconstructed
    assert len(doc.paragraphs) > 0


def test_pdf_to_powerpoint(sample_pdf, temp_dir):
    out_pptx = os.path.join(temp_dir, "output.pptx")
    res = pdf_to_powerpoint(sample_pdf, out_pptx)

    assert res["status"] == "success"
    assert os.path.exists(out_pptx)

    prs = Presentation(out_pptx)
    assert len(prs.slides) == 3
