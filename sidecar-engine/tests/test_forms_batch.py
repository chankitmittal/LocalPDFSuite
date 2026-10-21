"""
Unit tests for Batch Form Data to Excel/CSV extraction engine.
Tests:
- Schema union extraction across multiple filled AcroForms
- Formatted Excel (.xlsx) creation and content verification
- RFC 4180 CSV with UTF-8 BOM verification
"""

import os
import csv
import pytest
import openpyxl
from reportlab.pdfgen import canvas
from pypdf import PdfReader, PdfWriter
from pypdf.generic import (
    DictionaryObject,
    NameObject,
    TextStringObject,
    ArrayObject,
    NumberObject,
)

from core.pdf_forms_batch import (
    extract_form_fields_from_pdf,
    batch_extract_forms_to_spreadsheet,
)


def create_mock_form_pdf(output_path: str, fields: dict):
    """
    Creates a simple PDF with interactive AcroForm fields for testing.
    """
    # 1. Base blank PDF
    c = canvas.Canvas(output_path)
    c.drawString(100, 750, "Mock Form Document")
    c.save()

    # 2. Add AcroForm fields
    reader = PdfReader(output_path)
    writer = PdfWriter()
    page = writer.add_page(reader.pages[0])

    field_objs = []
    for f_name, f_val in fields.items():
        field_dict = DictionaryObject()
        field_dict[NameObject("/T")] = TextStringObject(f_name)
        field_dict[NameObject("/FT")] = NameObject("/Tx")
        field_dict[NameObject("/V")] = TextStringObject(str(f_val))
        field_dict[NameObject("/Rect")] = ArrayObject([
            NumberObject(100), NumberObject(700), NumberObject(300), NumberObject(720)
        ])
        field_dict[NameObject("/Subtype")] = NameObject("/Widget")
        writer_field = writer._add_object(field_dict)
        field_objs.append(writer_field)

    writer.set_need_appearances_writer()
    if not writer._root_object.get(NameObject("/AcroForm")):
        writer._root_object[NameObject("/AcroForm")] = DictionaryObject()
    writer._root_object[NameObject("/AcroForm")][NameObject("/Fields")] = ArrayObject(field_objs)

    with open(output_path, "wb") as f_out:
        writer.write(f_out)


def test_batch_form_data_extraction(tmp_path):
    # Form 1: Name, Email, Department
    pdf1 = str(tmp_path / "form1.pdf")
    create_mock_form_pdf(pdf1, {
        "FullName": "Alice Smith",
        "EmailAddress": "alice@company.local",
        "Department": "Legal",
    })

    # Form 2: Name, Email, Role (different schema from Form 1)
    pdf2 = str(tmp_path / "form2.pdf")
    create_mock_form_pdf(pdf2, {
        "FullName": "Bob Jones",
        "EmailAddress": "bob@company.local",
        "Role": "Compliance Officer",
    })

    # Form 3: Name, Phone (another variation)
    pdf3 = str(tmp_path / "form3.pdf")
    create_mock_form_pdf(pdf3, {
        "FullName": "Charlie Brown",
        "PhoneNumber": "555-0199",
    })

    # 1. Test single field extraction
    rec1 = extract_form_fields_from_pdf(pdf1)
    assert rec1["error"] is None
    assert rec1["field_count"] == 3
    assert rec1["fields"]["FullName"]["value"] == "Alice Smith"
    assert rec1["fields"]["EmailAddress"]["value"] == "alice@company.local"

    # 2. Test batch extraction
    out_dir = str(tmp_path / "batch_out")
    result = batch_extract_forms_to_spreadsheet([pdf1, pdf2, pdf3], out_dir, "test_export")

    assert result["status"] == "success"
    assert result["total_documents"] == 3
    # Union should contain FullName, EmailAddress, Department, Role, PhoneNumber
    expected_fields = {"FullName", "EmailAddress", "Department", "Role", "PhoneNumber"}
    assert set(s["field_name"] for s in result["schema"]) == expected_fields

    # 3. Verify .xlsx file
    xlsx_path = result["xlsx_path"]
    assert os.path.exists(xlsx_path)
    wb = openpyxl.load_workbook(xlsx_path)
    assert "Form Data" in wb.sheetnames
    assert "Field Schema & Stats" in wb.sheetnames

    ws = wb["Form Data"]
    # Row 1 is header, Rows 2-4 are data
    assert ws.max_row == 4
    # FullName appears in all 3 documents, so it should be the first data field after the metadata headers
    headers = [cell.value for cell in ws[1]]
    assert "Source File" in headers
    assert "FullName" in headers
    assert "EmailAddress" in headers

    # 4. Verify .csv file with UTF-8 BOM
    csv_path = result["csv_path"]
    assert os.path.exists(csv_path)

    # Check UTF-8 BOM
    with open(csv_path, "rb") as bf:
        raw_start = bf.read(3)
        assert raw_start == b"\xef\xbb\xbf", "CSV must contain UTF-8 BOM for Excel compatibility"

    # Check rows read correctly
    with open(csv_path, "r", encoding="utf-8-sig") as cf:
        reader = csv.reader(cf)
        rows = list(reader)
        assert len(rows) == 4  # Header + 3 records
